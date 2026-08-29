from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "ArtistCRM_trust_and_content_plan.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "20364D"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F4F6F9"
PALE_GOLD = "FFF7E6"
GOLD = "A66B00"
WHITE = "FFFFFF"
BLACK = "111827"
GREEN = "1F6B45"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_run_font(run, size=11, bold=None, color=BLACK, italic=None, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_para_format(p, before=0, after=6, line=1.25, keep_next=False):
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep_next


def add_body(doc, text="", bold_prefix=None, after=6, italic=False):
    p = doc.add_paragraph()
    set_para_format(p, after=after)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r, italic=italic)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    set_para_format(p, after=4)
    p.paragraph_format.left_indent = Inches(0.375 + 0.25 * level)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(0.375 + 0.25 * level))
    r = p.add_run(text)
    set_run_font(r)
    return p


def new_decimal_num_id(doc):
    numbering = doc.part.numbering_part.element
    num_ids = [
        int(node.get(qn("w:numId")))
        for node in numbering.findall(qn("w:num"))
    ]
    num_id = max(num_ids, default=0) + 1

    style_num_id = int(doc.styles["List Number"]._element.pPr.numPr.numId.val)
    base_num = next(
        node
        for node in numbering.findall(qn("w:num"))
        if int(node.get(qn("w:numId"))) == style_num_id
    )
    abstract_id = int(base_num.find(qn("w:abstractNumId")).get(qn("w:val")))

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    lvl_override = OxmlElement("w:lvlOverride")
    lvl_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    lvl_override.append(start_override)
    num.append(lvl_override)
    numbering.append(num)
    return num_id


def add_numbered_items(doc, items):
    num_id = new_decimal_num_id(doc)
    for text in items:
        p = doc.add_paragraph(style="List Number")
        set_para_format(p, after=4)
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)
        p_pr = p._p.get_or_add_pPr()
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num_id_el = OxmlElement("w:numId")
        num_id_el.set(qn("w:val"), str(num_id))
        num_pr.append(ilvl)
        num_pr.append(num_id_el)
        p_pr.append(num_pr)
        r = p.add_run(text)
        set_run_font(r)


def add_check(doc, text):
    return add_bullet(doc, f"[ ] {text}")


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p


def add_callout(doc, label, text, fill=LIGHT_BLUE, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    set_para_format(p, after=0)
    r1 = p.add_run(label + " ")
    set_run_font(r1, bold=True, color=accent)
    r2 = p.add_run(text)
    set_run_font(r2, color=BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_table(doc, headers, rows, widths_dxa, header_fill=LIGHT_BLUE, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths_dxa)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    set_row_cant_split(hdr)
    for idx, text in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        set_para_format(p, after=0, line=1.05)
        r = p.add_run(text)
        set_run_font(r, size=font_size, bold=True, color=INK)
    for row_data in rows:
        row = table.add_row()
        set_row_cant_split(row)
        for idx, text in enumerate(row_data):
            cell = row.cells[idx]
            p = cell.paragraphs[0]
            set_para_format(p, after=0, line=1.05)
            r = p.add_run(str(text))
            set_run_font(r, size=font_size, color=BLACK)
    set_table_geometry(table, widths_dxa)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Страница ")
    set_run_font(run, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for level, size, color, before, after in [
        (1, 16, BLUE, 18, 10),
        (2, 13, BLUE, 14, 7),
        (3, 12, DARK_BLUE, 10, 5),
    ]:
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for name in ["List Bullet", "List Bullet 2", "List Number"]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def configure_section(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_para_format(hp, after=0, line=1.0)
    rr = hp.add_run("ArtistCRM | Руководство по укреплению доверия")
    set_run_font(rr, size=9, color=MUTED)
    footer = section.footer
    add_page_number(footer.paragraphs[0])


def add_cover(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(110)
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("ПРАКТИЧЕСКОЕ РУКОВОДСТВО")
    set_run_font(r, size=10, bold=True, color=GOLD)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("Как укрепить доверие\nк ArtistCRM")
    set_run_font(r, size=30, bold=True, color=INK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(32)
    r = p.add_run("Пошаговый план на 8 недель, контент-план, площадки, сценарии и чек-листы")
    set_run_font(r, size=14, color=MUTED)

    add_callout(
        doc,
        "Цель:",
        "сделать ArtistCRM продуктом, за которым видны реальный основатель, профессиональный опыт, активные пользователи и проверяемые результаты.",
        fill=PALE_GOLD,
        accent=GOLD,
    )

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(55)
    r = p.add_run("Для основателя ArtistCRM | рабочая версия")
    set_run_font(r, size=10, color=MUTED, italic=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Август 2026")
    set_run_font(r, size=10, bold=True, color=INK)
    doc.add_page_break()


def build():
    doc = Document()
    configure_styles(doc)
    for section in doc.sections:
        configure_section(section)
    props = doc.core_properties
    props.title = "Как укрепить доверие к ArtistCRM"
    props.subject = "Пошаговый план и контент-план для основателя ArtistCRM"
    props.author = "ArtistCRM"
    props.keywords = "ArtistCRM, доверие, контент-план, продвижение, кейсы, отзывы"

    add_cover(doc)

    add_heading(doc, "Как пользоваться руководством", 1)
    add_body(doc, "Документ рассчитан на занятость 3–5 часов в неделю. Выполняйте этапы последовательно, отмечайте чек-листы и не переходите к масштабированию рекламы, пока не появились реальные подтверждения пользы продукта.")
    add_callout(doc, "Главный принцип:", "не заявлять, что ArtistCRM лучший, а регулярно показывать, кто его создаёт, для кого он работает и какие конкретные задачи решает.")

    add_heading(doc, "Что может сделать только основатель", 2)
    for item in [
        "лично рассказать историю продукта и показать профессиональный опыт артиста;",
        "поговорить с коллегами и услышать их настоящие формулировки проблем;",
        "провести демонстрации и вручную подключить первых пользователей;",
        "получить разрешения на публикацию отзывов, фотографий и кейсов;",
        "показать продукт в реальной работе и отвечать на неудобные вопросы;",
        "поддерживать отношения в профессиональных сообществах и просить рекомендации.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "Ожидаемый результат через 8 недель", 2)
    add_table(
        doc,
        ["Результат", "Целевой ориентир", "Как проверить"],
        [
            ("Содержательные диалоги", "20–30", "Есть записи проблем и возражений"),
            ("Личные демонстрации", "8–10", "Назначены следующие шаги"),
            ("Подключённые пользователи", "Не менее 5", "Внесены реальные заявки"),
            ("Активные пользователи", "Не менее 3", "Вернулись через 7 дней"),
            ("Публичные отзывы", "2–3", "Есть согласие и проверяемый автор"),
            ("Подробные кейсы", "1", "Опубликован на сайте"),
            ("Открытые демонстрации", "1", "Есть запись и вопросы аудитории"),
            ("Рекомендации коллег", "3–5", "Получены знакомства или диалоги"),
        ],
        [3000, 1800, 4560],
    )

    add_heading(doc, "Недельный рабочий ритм", 1)
    add_body(doc, "Повторяйте этот цикл каждую неделю. Он важнее ежедневной публикационной активности.")
    for item in [
        "5 личных диалогов с потенциальными пользователями;",
        "1–2 демонстрации ArtistCRM;",
        "1 интервью с пользователем или коллегой;",
        "1 содержательная публикация;",
        "1 короткое видео;",
        "1 улучшение страницы или продукта на основе обратной связи.",
    ]:
        add_check(doc, item)

    add_heading(doc, "План на 8 недель", 1)
    add_table(
        doc,
        ["Неделя", "Главная задача", "Итог недели"],
        [
            ("1", "Записать историю создания ArtistCRM", "Видео и текст об основателе"),
            ("2", "Провести 3–5 проблемных интервью", "Язык клиентов и список возражений"),
            ("3", "Ручное подключение двух пилотов", "Реальные сценарии использования"),
            ("4", "Получить два честных отзыва", "Первые социальные доказательства"),
            ("5", "Снять продукт в реальной работе", "Демонстрационный ролик"),
            ("6", "Оформить первый подробный кейс", "Проверяемое доказательство пользы"),
            ("7", "Провести открытый эфир", "Запись, вопросы и новые диалоги"),
            ("8", "Запустить рекомендации", "Новые знакомства от пользователей"),
        ],
        [1100, 4300, 3960],
    )

    add_heading(doc, "Этап 1. История основателя", 1)
    add_body(doc, "Запишите видео на 2–4 минуты. Небольшая естественная неидеальность убедительнее постановочной рекламы.")
    add_heading(doc, "Сценарий видео", 2)
    add_numbered_items(doc, [
        "Кто вы и сколько лет работаете артистом.",
        "Как раньше учитывали заявки и даты.",
        "Какие ошибки и неудобства возникали.",
        "Почему одного календаря оказалось недостаточно.",
        "Почему вы начали создавать ArtistCRM.",
        "Как сами используете систему.",
        "Кому хотите помочь.",
        "Почему лично отвечаете первым пользователям.",
    ])
    add_heading(doc, "Материалы, которые нужно собрать", 2)
    for item in [
        "актуальная фотография основателя;",
        "краткая профессиональная биография;",
        "ссылки, подтверждающие опыт артиста;",
        "2–4 фотографии с выступлений;",
        "год начала профессиональной деятельности;",
        "короткое объяснение, почему ArtistCRM существует.",
    ]:
        add_check(doc, item)

    add_heading(doc, "Этап 2. Проблемные интервью", 1)
    add_body(doc, "Поговорите минимум с пятью коллегами. Не начинайте с презентации ArtistCRM: сначала изучите их текущий способ работы.")
    add_heading(doc, "Вопросы интервью", 2)
    questions = [
        "Откуда обычно приходят заявки?",
        "Где они сейчас записываются?",
        "Случалось ли забыть ответить или перезвонить?",
        "Как проверяется свободная дата?",
        "Где учитываются задатки и остатки оплаты?",
        "Какие действия приходится постоянно держать в голове?",
        "Что раздражает в таблицах и календаре?",
        "Как должна выглядеть система, которой вы действительно будете пользоваться?",
        "Что вызывает недоверие к новой CRM?",
        "Что должно произойти, чтобы вы согласились попробовать?",
    ]
    add_numbered_items(doc, questions)
    add_callout(doc, "Важно:", "перед записью разговора получите согласие. Для публикации цитаты, имени, фотографии или ссылки согласуйте отдельное разрешение.", fill=PALE_GOLD, accent=GOLD)

    add_heading(doc, "Карточка одного интервью", 2)
    add_table(
        doc,
        ["Поле", "Запись"],
        [
            ("Профессия и тип работы", ""),
            ("Текущий способ учёта", ""),
            ("Главная проблема", ""),
            ("Точная формулировка пользователя", ""),
            ("Главное возражение", ""),
            ("Что готов попробовать", ""),
            ("Следующий шаг и дата", ""),
        ],
        [2700, 6660],
    )

    add_heading(doc, "Этап 3. Ручное подключение пилота", 1)
    add_body(doc, "Не ограничивайтесь выдачей ссылки. Первое доверие возникает, когда основатель лично помогает получить результат.")
    for item in [
        "заполнить профиль;",
        "создать 2–3 услуги;",
        "внести минимум две реальные заявки;",
        "назначить следующее действие;",
        "отметить задаток или ожидаемую оплату;",
        "создать мероприятие;",
        "включить уведомления и установить PWA;",
        "при необходимости подключить Google Calendar;",
        "назначить повторный разговор через семь дней.",
    ]:
        add_check(doc, item)
    add_callout(doc, "Правильный вопрос после подключения:", "Что сейчас кажется непонятным или вызывает сомнения? Не заменяйте его вопросом «Вам всё понравилось?»")

    add_heading(doc, "Этап 4. Отзывы", 1)
    add_body(doc, "Просите отзыв через 7–10 дней, когда пользователь успел выполнить реальные действия.")
    add_heading(doc, "Шаблон запроса", 2)
    add_callout(doc, "Сообщение:", "Привет! Спасибо, что тестируешь ArtistCRM. Можешь честно написать в двух-трёх предложениях: как раньше учитывал заявки, что оказалось полезным и что пока неудобно? Если разрешишь, я размещу отзыв на сайте с твоим именем и профессией. Перед публикацией обязательно покажу итоговый вариант.", fill=LIGHT_GRAY, accent=DARK_BLUE)
    add_heading(doc, "Проверка отзыва перед публикацией", 2)
    for item in [
        "есть имя или согласованный псевдоним;",
        "указана профессия;",
        "есть фотография или ссылка на публичную страницу — по желанию;",
        "описана прежняя проблема;",
        "названа конкретная полезная функция;",
        "нет неподтверждённых результатов;",
        "получено явное разрешение на публикацию;",
        "итоговая версия повторно показана автору.",
    ]:
        add_check(doc, item)

    add_heading(doc, "Этап 5. Демонстрация реальной работы", 1)
    add_body(doc, "Запишите ролик «Как я сам веду заявку в ArtistCRM». Показывайте один завершённый сценарий, а не все функции одновременно.")
    add_numbered_items(doc, [
        "Пришло сообщение клиента.",
        "Создаётся заявка.",
        "Указывается дата.",
        "Назначается следующий контакт.",
        "Добавляется задаток.",
        "Заявка превращается в мероприятие.",
        "Проверяется занятость даты.",
    ])
    add_callout(doc, "Безопасность:", "используйте демонстрационные или обезличенные данные. Не показывайте реальные телефоны, фамилии, адреса, переписку и финансовые сведения клиентов.", fill=PALE_GOLD, accent=GOLD)

    add_heading(doc, "Этап 6. Первый кейс", 1)
    add_body(doc, "Проведите 20-минутное интервью с активным пользователем и оформите доказательство результата без преувеличений.")
    add_heading(doc, "Структура кейса", 2)
    add_numbered_items(doc, [
        "Профессия пользователя и рабочий контекст.",
        "Как заявки учитывались раньше.",
        "Какая проблема возникала чаще всего.",
        "Почему пользователь решил попробовать ArtistCRM.",
        "Как проходило подключение.",
        "Какие функции используются регулярно.",
        "Что стало проще через 1–3 недели.",
        "Что в продукте ещё требуется улучшить.",
        "Готов ли пользователь продолжать работу.",
    ])
    add_body(doc, "Допустимые первые результаты: все обращения находятся в одном месте; видно, кому нужно перезвонить; задатки перестали храниться в голове; свободная дата проверяется быстрее; снизилось беспокойство, что что-то забыто.")

    add_heading(doc, "Этап 7. Открытая демонстрация", 1)
    add_body(doc, "Тема: «Как артисту или частному специалисту не терять заявки, задатки и перезвоны».")
    add_table(
        doc,
        ["Время", "Блок", "Содержание"],
        [
            ("5 минут", "История", "Почему возник ArtistCRM"),
            ("10 минут", "Проблемы", "Типичные ошибки учёта заявок"),
            ("15 минут", "Демонстрация", "Один рабочий сценарий целиком"),
            ("10 минут", "Вопросы", "Ответы без ухода от сложных тем"),
            ("2 минуты", "CTA", "Приглашение на личную настройку"),
        ],
        [1500, 2200, 5660],
    )
    add_body(doc, "После эфира подготовьте: полную запись, 5–7 коротких фрагментов, статью, ответы для FAQ и несколько постов.")

    add_heading(doc, "Этап 8. Рекомендации", 1)
    add_callout(doc, "Шаблон:", "Если ArtistCRM оказался полезен, можешь познакомить меня с одним коллегой, которому тоже приходится вести заявки и задатки? Я лично помогу ему всё настроить. За активного приглашённого пользователя добавлю тебе месяц тарифа «Бизнес».", fill=LIGHT_GRAY, accent=DARK_BLUE)
    add_body(doc, "Вознаграждение начисляйте за активированного пользователя, а не за пустую регистрацию.")

    doc.add_page_break()
    add_heading(doc, "Где публиковать материалы", 1)
    add_table(
        doc,
        ["Площадка", "Роль", "Что размещать"],
        [
            ("Профессиональные чаты", "Первые пользователи", "Личная история, комикс, демонстрация, приглашение в пилот"),
            ("Личная страница VK", "Доверие к основателю", "Опыт артиста, наблюдения, видео, отзывы, изменения продукта"),
            ("Сообщество ArtistCRM VK", "Публичная витрина", "Закреплённая история, тарифы, ответы, отзывы, обновления"),
            ("Сайт ArtistCRM", "Главный архив доказательств", "Кейсы, инструкции, безопасность, история основателя"),
            ("Telegram", "Живые заметки", "Обновления, вопросы, короткие видео, приглашения на тест"),
            ("VK Видео / YouTube", "Демонстрация", "Короткие сценарии, обзоры, записи эфиров"),
            ("Дзен", "Накопительный охват", "Проблемные статьи после появления 5–7 материалов"),
            ("VC.ru", "История продукта", "Редкие подробные материалы с опытом, цифрами и ошибками"),
        ],
        [2300, 2550, 4510],
        font_size=9,
    )

    add_heading(doc, "Приоритет площадок", 2)
    add_body(doc, "Обязательные на старте: личная страница VK, профессиональные чаты, сообщество ArtistCRM во VK и сайт. Telegram и видео — дополнительные. Дзен и VC.ru подключайте после появления сильных кейсов.")
    add_callout(doc, "Принцип распространения:", "сайт хранит доказательства, личная страница создаёт доверие, профессиональные чаты приводят первых пользователей, а сообщество ArtistCRM показывает, что продукт развивается.")

    add_heading(doc, "Как распространять один материал", 2)
    add_numbered_items(doc, [
        "Полная версия — на artistcrm.ru.",
        "Личная история и ссылка — на странице основателя VK.",
        "Краткая версия — в сообществе ArtistCRM.",
        "Одно изображение и 3–4 абзаца — в тематическом чате.",
        "Короткий ролик — в VK Видео и YouTube Shorts.",
        "Небольшое наблюдение — в Telegram.",
        "После накопления фактуры — обобщающая статья в Дзене или на VC.ru.",
    ])

    doc.add_page_break()
    add_heading(doc, "Контент-план на 4 недели", 1)
    add_table(
        doc,
        ["Неделя", "Основной пост", "Короткое видео", "CTA"],
        [
            ("1", "Почему практикующий артист начал делать CRM", "Как теряется заявка между мессенджерами", "Напишите «CRM» — покажу систему"),
            ("2", "Пять вещей, которые опасно держать в голове", "Создание заявки и перезвона за минуту", "Приглашение на личную настройку"),
            ("3", "Почему календарь не заменяет учёт клиентов", "Задаток и остаток оплаты", "Получить демонстрацию"),
            ("4", "Первый отзыв или мини-кейс", "Что изменилось после разговора с коллегой", "Попробовать пилот на два месяца"),
        ],
        [1000, 3300, 2900, 2160],
        font_size=8.7,
    )

    add_heading(doc, "Формула хорошего поста", 2)
    add_numbered_items(doc, [
        "Реальная ситуация.",
        "Почему она опасна или неудобна.",
        "Как вы решали её раньше.",
        "Как решаете сейчас.",
        "Скриншот или короткое видео.",
        "Один понятный призыв к действию.",
    ])
    add_callout(doc, "Пример CTA:", "Я набираю несколько коллег на бесплатное двухмесячное тестирование с личной настройкой. Если хотите посмотреть — напишите мне слово «CRM».", fill=LIGHT_GRAY, accent=DARK_BLUE)

    add_heading(doc, "Дополнительные темы", 2)
    for item in [
        "как проверить свободную дату;",
        "что записывать после первого звонка клиента;",
        "как не забывать о неоплаченном остатке;",
        "CRM против Excel;",
        "зачем артисту история клиента;",
        "как принимать заявки с Tilda;",
        "ошибки при ведении заявок в мессенджерах;",
        "разбор рабочего сценария без персональных данных.",
    ]:
        add_bullet(doc, item)

    doc.add_page_break()
    add_heading(doc, "Пример недельного графика", 1)
    add_table(
        doc,
        ["День", "Действие", "Площадка"],
        [
            ("Вторник", "Полезный пост", "Личная страница VK"),
            ("Среда", "Адаптированная версия", "Сообщество ArtistCRM"),
            ("Четверг или пятница", "Сообщение в одном тематическом чате", "Профессиональное сообщество"),
            ("Суббота", "Короткое видео", "VK Видео / YouTube"),
            ("Раз в месяц", "Большой кейс или статья", "Сайт ArtistCRM"),
            ("Раз в месяц", "Открытая демонстрация", "VK / Telegram / видеоплощадка"),
        ],
        [2100, 4050, 3210],
    )

    add_heading(doc, "Материалы доверия для сайта", 1)
    add_body(doc, "Собирайте фактуру постепенно. После появления материалов их можно оформить в отдельные страницы и блоки сайта.")
    for item in [
        "фотография и история основателя;",
        "подтверждённый профессиональный опыт;",
        "реальные экраны продукта;",
        "публичная цена без скрытых условий;",
        "понятный бесплатный тариф;",
        "информация о защите и удалении данных;",
        "контакт и ожидаемое время ответа поддержки;",
        "дата последнего обновления продукта;",
        "открытый список улучшений;",
        "отзывы с проверяемыми авторами;",
        "честные кейсы и ответы на неудобные вопросы;",
        "понятные реквизиты, документы и условия использования.",
    ]:
        add_check(doc, item)

    add_heading(doc, "Чего не делать", 1)
    for item in [
        "не использовать вымышленные отзывы и фотографии из фотостоков вместо пользователей;",
        "не публиковать неподтверждённые цифры и обещания роста дохода;",
        "не ставить фальшивые таймеры и искусственный дефицит;",
        "не заявлять «№1 CRM для артистов» без проверяемого основания;",
        "не рассылать одинаковый рекламный текст во все сообщества одновременно;",
        "не показывать персональные данные клиентов в видео и скриншотах;",
        "не создавать десятки одинаковых SEO-страниц без реальной пользы.",
    ]:
        add_bullet(doc, item)
    add_callout(doc, "Тон коммуникации:", "спокойная честность молодого продукта убедительнее громких обещаний. Показывайте не только успехи, но и то, что изменили после обратной связи.", fill=PALE_GOLD, accent=GOLD)

    add_heading(doc, "Еженедельный контроль", 1)
    add_table(
        doc,
        ["Показатель", "За неделю", "Комментарий / следующий шаг"],
        [
            ("Личные диалоги", "", ""),
            ("Демонстрации", "", ""),
            ("Новые пилоты", "", ""),
            ("Активированные пользователи", "", ""),
            ("Полученные отзывы", "", ""),
            ("Опубликованные материалы", "", ""),
            ("Рекомендации", "", ""),
            ("Главное возражение недели", "", ""),
            ("Что нужно изменить", "", ""),
        ],
        [3100, 1700, 4560],
    )

    add_heading(doc, "Финальная проверка готовности доверительной базы", 1)
    for item in [
        "На сайте понятно, кто создал ArtistCRM и почему.",
        "Есть минимум два согласованных отзыва.",
        "Опубликован минимум один подробный кейс.",
        "Есть видео полного рабочего сценария.",
        "Пользователь легко находит тарифы, безопасность и поддержку.",
        "Все заявления можно подтвердить.",
        "Для каждой профессии используется свой язык и свой пример.",
        "Из рекламы можно перейти на релевантную посадочную.",
        "Регистрация и первая активация проверены вручную.",
        "Известно, какой источник приводит активных пользователей.",
    ]:
        add_check(doc, item)

    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(180)
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("Следующее действие")
    set_run_font(r, size=25, bold=True, color=INK)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(28)
    r = p.add_run("Не начинайте с ещё одной рекламной публикации")
    set_run_font(r, size=13, color=MUTED, italic=True)
    add_callout(doc, "Начните сегодня:", "назначьте три разговора с коллегами и запишите черновик двухминутной истории о том, почему вы создали ArtistCRM.", fill=LIGHT_BLUE, accent=BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(28)
    r = p.add_run("Живой опыт основателя — главное конкурентное преимущество ArtistCRM на старте.")
    set_run_font(r, size=11, bold=True, color=DARK_BLUE)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
